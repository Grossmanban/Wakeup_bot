"""Импорт заданий из Excel: разбор → предпросмотр с ошибками → подтверждение (ТЗ 5)."""

import asyncio
import io
from dataclasses import dataclass
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func, select
from sqlalchemy.dialects.postgresql import insert
from sqlalchemy.ext.asyncio import AsyncSession

from app.errors import AppError, Conflict, NotFound
from app.models import ImportBatch, Product, ProductStatus, Task, TaskStatus, TaskType
from app.permissions import Perm
from app.services import audit
from app.services.access import Actor
from app.services.audit import Category
from app.timeutil import now_utc

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20000

# Поле задания → варианты заголовка в Excel (без учёта регистра и пробелов по краям).
HEADERS: dict[str, tuple[str, ...]] = {
    "external_id": ("id", "id задания", "уникальный id", "ид", "номер задания", "task id"),
    "product": ("продукт", "product"),
    "type": ("тип", "тип задания", "type"),
    "link": ("ссылка", "url", "link"),
    "expected_text": ("ожидаемый текст", "текст", "expected text"),
    "rating": ("рейтинг", "оценка", "rating"),
    "instructions": ("инструкция", "инструкции", "instructions"),
    "comment": ("комментарий", "комментарии", "comment"),
    "materials": ("материалы", "ссылки на материалы", "materials"),
    "gender": ("пол", "гендер", "gender"),
    "units": ("единиц", "единицы", "units"),
}
REQUIRABLE = ("link", "expected_text", "rating", "instructions", "comment", "materials", "gender")
FIELD_TITLES = {
    "external_id": "ID задания",
    "link": "Ссылка",
    "expected_text": "Ожидаемый текст",
    "rating": "Рейтинг",
    "instructions": "Инструкция",
    "comment": "Комментарий",
    "materials": "Материалы",
    "gender": "Пол",
}
GENDERS = {
    "м": "m",
    "муж": "m",
    "мужской": "m",
    "m": "m",
    "male": "m",
    "ж": "f",
    "жен": "f",
    "женский": "f",
    "f": "f",
    "female": "f",
    "": None,
    "любой": None,
    "-": None,
    "—": None,
    "any": None,
}


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


@dataclass
class Parsed:
    columns: dict[str, int]  # поле задания → номер колонки
    extra: list[tuple[int, str]]  # прочие колонки: (номер, заголовок)
    rows: list[list[str]]


def parse_xlsx(content: bytes) -> Parsed:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001 — openpyxl бросает разные исключения на битый файл
        raise AppError("bad_file", "Не удалось прочитать файл. Нужен формат .xlsx") from e
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise AppError("bad_file", "Файл пустой")
        names = [_cell(h) for h in header]
        columns: dict[str, int] = {}
        extra: list[tuple[int, str]] = []
        for i, h in enumerate(names):
            field = next((f for f, al in HEADERS.items() if h.lower() in al), None)
            if field and field not in columns:
                columns[field] = i
            elif h:
                extra.append((i, h))
        if "external_id" not in columns:
            raise AppError("bad_file", "Не найдена колонка с ID задания (например «ID задания»)")
        data = []
        for r in rows:
            if len(data) >= MAX_ROWS:
                raise AppError("too_many_rows", f"Слишком много строк: не больше {MAX_ROWS}")
            cells = [_cell(v) for v in r]
            data.append(cells + [""] * max(0, len(names) - len(cells)))
        return Parsed(columns, extra, data)
    finally:
        wb.close()


async def preview(
    session: AsyncSession,
    actor: Actor,
    *,
    file_name: str,
    content: bytes,
    product_id: int | None,
) -> ImportBatch:
    actor.require(Perm.TASKS_IMPORT)
    if len(content) > MAX_FILE_BYTES:
        raise AppError("too_large", "Файл больше 10 МБ")
    parsed = await asyncio.to_thread(parse_xlsx, content)
    mapping = parsed.columns

    products = {
        p.id: p
        for p in (
            await session.execute(select(Product).where(Product.status != ProductStatus.DELETED))
        ).scalars()
    }
    by_name = {p.name.lower(): p for p in products.values()}
    default_product = None
    if product_id is not None:
        default_product = products.get(product_id)
        if default_product is None:
            raise NotFound("Продукт не найден")
    if "product" not in mapping and default_product is None:
        raise AppError("no_product", "В файле нет колонки «Продукт» — выберите продукт для импорта")

    types: dict[int, list[TaskType]] = {}
    for t in (
        await session.execute(select(TaskType).where(TaskType.is_active.is_(True)))
    ).scalars():
        types.setdefault(t.product_id, []).append(t)

    def get(row: list[str], field: str) -> str:
        i = mapping.get(field)
        return row[i] if i is not None and i < len(row) else ""

    out: list[dict[str, Any]] = []
    seen: set[tuple[int, str]] = set()
    for n, row in enumerate(parsed.rows, start=2):
        if not any(row):
            out.append({"row": n, "status": "skipped", "errors": []})
            continue
        errors: list[str] = []
        ext = get(row, "external_id")
        if not ext:
            errors.append("Нет ID задания")
        elif len(ext) > 128:
            errors.append("ID задания длиннее 128 символов")

        product = default_product
        pname = get(row, "product")
        if pname:
            product = by_name.get(pname.lower())
            if product is None:
                errors.append(f"Неизвестный продукт «{pname}»")
        if product is not None and product.status == ProductStatus.ARCHIVED:
            errors.append(f"Продукт «{product.name}» в архиве")

        task_type = None
        tname = get(row, "type")
        if product is not None:
            ptypes = types.get(product.id, [])
            if tname:
                key = " ".join(tname.split()).lower()
                task_type = next(
                    (
                        t
                        for t in ptypes
                        if key == t.name.lower() or key in {a.lower() for a in t.excel_names}
                    ),
                    None,
                )
                if task_type is None:
                    errors.append(f"Неизвестный тип задания «{tname}»")
            elif len(ptypes) == 1:
                task_type = ptypes[0]
            elif ptypes:
                errors.append("Не указан тип задания")

        gender_raw = get(row, "gender").lower()
        if gender_raw not in GENDERS:
            errors.append(f"Непонятное значение пола «{get(row, 'gender')}»")
        gender = GENDERS.get(gender_raw)

        units = task_type.units if task_type else 1
        units_raw = get(row, "units")
        if units_raw:
            if units_raw.isdigit() and 0 < int(units_raw) <= 1000:
                units = int(units_raw)
            else:
                errors.append(f"Некорректное число единиц «{units_raw}»")

        link = get(row, "link")
        if link and not link.lower().startswith(("http://", "https://")):
            errors.append("Ссылка должна начинаться с http:// или https://")

        expected = get(row, "expected_text") or (product.expected_text_default if product else "")
        values = {
            "link": link,
            "expected_text": expected,
            "rating": get(row, "rating"),
            "instructions": get(row, "instructions"),
            "comment": get(row, "comment"),
            "materials": get(row, "materials"),
            "gender": gender_raw,
        }
        if product is not None:
            for f in product.required_fields:
                if f in values and not values[f]:
                    errors.append(f"Не заполнено обязательное поле «{FIELD_TITLES.get(f, f)}»")

        status = "error" if errors else "ok"
        if not errors:
            key = (product.id, ext)
            if key in seen:
                status = "duplicate"
                errors.append("Повтор ID в этом файле")
            seen.add(key)

        out.append(
            {
                "row": n,
                "status": status,
                "errors": errors,
                "product_id": product.id if product else None,
                "product": product.name if product else pname,
                "external_id": ext,
                "task_type_id": task_type.id if task_type else None,
                "type": task_type.name if task_type else tname,
                "link": link or None,
                "expected_text": expected or None,
                "rating": values["rating"][:32] or None,
                "instructions": values["instructions"] or None,
                "comment": values["comment"] or None,
                "materials": [m.strip() for m in values["materials"].split() if m.strip()],
                "gender": gender,
                "units": units,
                "extra": {name: row[i] for i, name in parsed.extra if i < len(row) and row[i]},
            }
        )

    # Уже загруженные ранее задания — тоже дубликаты.
    ok_keys = [(r["product_id"], r["external_id"]) for r in out if r["status"] == "ok"]
    existing: set[tuple[int, str]] = set()
    for pid in {k[0] for k in ok_keys}:
        ids = [k[1] for k in ok_keys if k[0] == pid]
        for chunk in range(0, len(ids), 1000):
            existing |= {
                (pid, e)
                for e in (
                    await session.execute(
                        select(Task.external_id).where(
                            Task.product_id == pid, Task.external_id.in_(ids[chunk : chunk + 1000])
                        )
                    )
                ).scalars()
            }
    for r in out:
        if r["status"] == "ok" and (r["product_id"], r["external_id"]) in existing:
            r["status"] = "duplicate"
            r["errors"].append("Задание с таким ID уже загружено")

    summary = {
        "total": len(out),
        "ok": sum(r["status"] == "ok" for r in out),
        "errors": sum(r["status"] == "error" for r in out),
        "duplicates": sum(r["status"] == "duplicate" for r in out),
        "skipped": sum(r["status"] == "skipped" for r in out),
        "columns": sorted(mapping),
        "extra_columns": [name for _, name in parsed.extra],
    }
    batch = ImportBatch(
        product_id=product_id,
        file_name=file_name[:255],
        uploaded_by_id=actor.id,
        rows=out,
        summary=summary,
    )
    session.add(batch)
    await session.commit()
    return batch


async def commit(session: AsyncSession, actor: Actor, batch_id: int) -> dict[str, Any]:
    """Загрузить строки без ошибок. Повторное подтверждение возвращает тот же результат."""
    actor.require(Perm.TASKS_IMPORT)
    batch = await session.get(ImportBatch, batch_id, with_for_update=True)
    if batch is None or (batch.uploaded_by_id != actor.id and not actor.is_owner):
        raise NotFound("Загрузка не найдена")
    if batch.status == "committed":
        return batch.result or {}
    if batch.status != "preview":
        raise Conflict("bad_state", "Загрузка уже отменена")

    added = 0
    conflicts = 0
    per_product: dict[int, int] = {}
    for r in batch.rows:
        if r["status"] != "ok":
            continue
        res = await session.execute(
            insert(Task)
            .values(
                product_id=r["product_id"],
                task_type_id=r["task_type_id"],
                external_id=r["external_id"],
                link=r["link"],
                expected_text=r["expected_text"],
                rating=r["rating"],
                instructions=r["instructions"],
                comment=r["comment"],
                materials=r["materials"],
                gender=r["gender"],
                extra=r["extra"],
                units=r["units"],
                import_batch_id=batch.id,
                status=TaskStatus.POOL,
            )
            .on_conflict_do_nothing(constraint="uq_tasks_product_external")
        )
        if res.rowcount:
            added += 1
            per_product[r["product_id"]] = per_product.get(r["product_id"], 0) + 1
        else:
            conflicts += 1

    available = {}
    for pid in per_product:
        available[str(pid)] = await session.scalar(
            select(func.count()).where(Task.product_id == pid, Task.status == TaskStatus.POOL)
        )
    s = batch.summary
    result = {
        "added": added,
        "skipped": s["skipped"],
        "duplicates": s["duplicates"] + conflicts,
        "errors": s["errors"],
        "available": available,
    }
    batch.status = "committed"
    batch.committed_at = now_utc()
    batch.result = result
    audit.record(
        session,
        category=Category.TASKS,
        action="tasks.imported",
        actor_id=actor.id,
        object_type="import_batch",
        object_id=batch.id,
        new={"file": batch.file_name, **result},
    )
    await session.commit()
    return result


async def discard(session: AsyncSession, actor: Actor, batch_id: int) -> None:
    actor.require(Perm.TASKS_IMPORT)
    batch = await session.get(ImportBatch, batch_id, with_for_update=True)
    if batch is None or (batch.uploaded_by_id != actor.id and not actor.is_owner):
        raise NotFound("Загрузка не найдена")
    if batch.status == "preview":
        batch.status = "discarded"
        await session.commit()
