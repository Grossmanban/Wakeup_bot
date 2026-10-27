"""Этап 9: статистика, производительность с учётом качества, экспорт."""

import io
from datetime import datetime, timedelta

from openpyxl import load_workbook

from app import timeutil
from app.db import get_sessionmaker
from app.models import Task, TaskStatus
from app.services.stats import merged_hours
from tests.helpers import OWNER_TG, auth, grant, make_active, owner_ready, ready_product

DAY = timeutil.msk_today() - timedelta(days=1)
BASE = datetime(DAY.year, DAY.month, DAY.day, 10, 0, tzinfo=timeutil.MSK)
Q = f"start={DAY}&end={DAY}"


async def _tasks(user_id: int, product_id: int, spec: list[tuple[int, int, TaskStatus]]):
    """spec: (начало в минутах от 10:00, длительность в минутах, статус)."""
    async with get_sessionmaker()() as s:
        for i, (start, minutes, status) in enumerate(spec):
            found = BASE + timedelta(minutes=start)
            s.add(
                Task(
                    product_id=product_id,
                    external_id=f"S{user_id}-{i}-{start}",
                    status=status,
                    assignee_id=user_id,
                    assigned_at=found,
                    found_at=found,
                    completed_at=found + timedelta(minutes=minutes),
                    units=1,
                )
            )
        await s.commit()


def _user(data, uid):
    return next(u for u in data["users"] if u["user_id"] == uid)


def test_merged_hours_does_not_double_count():
    a = BASE
    assert (
        merged_hours(
            [(a, a + timedelta(hours=1)), (a + timedelta(minutes=30), a + timedelta(minutes=90))]
        )
        == 1.5
    )
    assert (
        merged_hours(
            [(a, a + timedelta(hours=1)), (a + timedelta(hours=2), a + timedelta(hours=3))]
        )
        == 2.0
    )


async def test_ten_per_hour_with_half_quality_gives_five(client):
    await owner_ready(client)
    emp = await make_active(client, 140)
    p = await ready_product(client, [emp], n_tasks=0)
    # 10 заданий по 6 минут подряд = 1 час; 5 одобрено, 5 отклонено.
    spec = [(i * 6, 6, TaskStatus.APPROVED if i < 5 else TaskStatus.REJECTED) for i in range(10)]
    await _tasks(emp, p["id"], spec)
    d = (await client.get(f"/api/stats?{Q}", headers=auth(OWNER_TG))).json()
    u = _user(d, emp)
    assert u["units"] == 10 and u["hours"] == 1.0
    assert u["units_per_hour"] == 10.0
    assert u["quality"] == 50.0
    assert u["forecast_paid_units_per_hour"] == 5.0
    assert u["sample_units"] == 10 and u["small_sample"] is True


async def test_eight_per_hour_with_95_quality_gives_7_6(client):
    await owner_ready(client)
    emp = await make_active(client, 141)
    p = await ready_product(client, [emp], n_tasks=0)
    # 20 заданий по 7,5 минуты = 2,5 часа → 8 ед./ч; 19 одобрено, 1 отклонено → 95%.
    spec = [
        (int(i * 7.5), 7, TaskStatus.APPROVED if i < 19 else TaskStatus.REJECTED) for i in range(20)
    ]
    await _tasks(emp, p["id"], spec)
    async with get_sessionmaker()() as s:  # точные 7,5 минуты
        from sqlalchemy import update

        for i in range(20):
            found = BASE + timedelta(minutes=i * 7.5)
            await s.execute(
                update(Task)
                .where(Task.assignee_id == emp, Task.external_id == f"S{emp}-{i}-{int(i * 7.5)}")
                .values(found_at=found, completed_at=found + timedelta(minutes=7.5))
            )
        await s.commit()
    u = _user((await client.get(f"/api/stats?{Q}", headers=auth(OWNER_TG))).json(), emp)
    assert u["units_per_hour"] == 8.0 and u["quality"] == 95.0
    assert u["forecast_paid_units_per_hour"] == 7.6
    assert u["small_sample"] is False


async def test_actual_paid_rate_pending_and_overlap(client):
    await owner_ready(client)
    emp = await make_active(client, 142)
    p = await ready_product(client, [emp], n_tasks=0)
    await _tasks(
        emp,
        p["id"],
        [
            (0, 60, TaskStatus.APPROVED),  # 10:00–11:00
            (30, 60, TaskStatus.APPROVED),  # 10:30–11:30 — пересекается
            (120, 30, TaskStatus.MODERATION),  # ожидает модерации
        ],
    )
    u = _user((await client.get(f"/api/stats?{Q}", headers=auth(OWNER_TG))).json(), emp)
    assert u["hours"] == 2.0  # 1,5 ч объединённых + 0,5 ч, а не 2,5
    assert u["units_per_hour"] == 1.5
    assert u["paid_units_per_hour"] == round(2 / 1.5, 2)  # только одобренные и их время
    assert u["pending_units"] == 1
    assert u["quality"] == 100.0 and u["sample_units"] == 2  # ожидающие не учитываются


async def test_insufficient_time_and_no_decisions(client):
    await owner_ready(client)
    emp = await make_active(client, 143)
    p = await ready_product(client, [emp], n_tasks=0)
    await _tasks(emp, p["id"], [(0, 5, TaskStatus.MODERATION)])
    u = _user((await client.get(f"/api/stats?{Q}", headers=auth(OWNER_TG))).json(), emp)
    assert u["units_per_hour"] is None and u["insufficient_time"] is True
    assert u["quality"] is None  # «—»: решений ещё нет


async def test_permissions_my_stats_and_export(client):
    await owner_ready(client)
    emp = await make_active(client, 144)
    analyst = await make_active(client, 145, role="admin")
    p = await ready_product(client, [emp], n_tasks=0)
    await _tasks(emp, p["id"], [(0, 30, TaskStatus.APPROVED), (40, 30, TaskStatus.APPROVED)])
    assert (await client.get(f"/api/stats?{Q}", headers=auth(144))).status_code == 403
    mine = (await client.get(f"/api/my/stats?{Q}", headers=auth(144))).json()
    assert [u["user_id"] for u in mine["users"]] == [emp] and mine["users"][0]["units"] == 2

    await grant(client, analyst, "stats.view")
    d = (await client.get(f"/api/stats?{Q}&product_ids={p['id']}", headers=auth(145))).json()
    assert d["products"][0]["units"] == 2
    assert d["hourly"][10]["units"] == 1 and d["hourly"][11]["units"] == 1
    assert "Прогноз" in d["formulas"]["forecast_paid_per_hour"]
    tasks = (await client.get(f"/api/stats/tasks?{Q}", headers=auth(145))).json()
    assert len(tasks) == 2 and tasks[0]["minutes"] == 30.0

    assert (await client.get(f"/api/stats/export?{Q}", headers=auth(145))).status_code == 403
    await grant(client, analyst, "stats.view", "export")
    r = await client.get(f"/api/stats/export?{Q}&kind=tasks&format=xlsx", headers=auth(145))
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb.active.iter_rows(values_only=True))
    assert rows[0][2] == "Продукт" and len(rows) == 3
    r = await client.get(f"/api/stats/export?{Q}&kind=users&format=csv", headers=auth(145))
    text = r.content.decode("utf-8-sig")
    assert text.splitlines()[0].startswith("Сотрудник;Выполнено единиц")


async def test_presets(client):
    from app.services.stats import period_preset

    wed = datetime(2026, 10, 28).date()
    assert period_preset("week", wed) == (datetime(2026, 10, 26).date(), wed)
    assert period_preset("last_week", wed) == (
        datetime(2026, 10, 19).date(),
        datetime(2026, 10, 25).date(),
    )
    assert period_preset("last_month", wed) == (
        datetime(2026, 9, 1).date(),
        datetime(2026, 9, 30).date(),
    )


async def test_download_link(client):
    await owner_ready(client)
    emp = await make_active(client, 146)
    p = await ready_product(client, [emp], n_tasks=0)
    await _tasks(emp, p["id"], [(0, 30, TaskStatus.APPROVED)])
    r = await client.post(
        "/api/stats/export-link",
        json={"kind": "tasks", "format": "csv", "start": str(DAY), "end": str(DAY)},
        headers=auth(OWNER_TG),
    )
    assert r.status_code == 200, r.text
    url = r.json()["url"]
    got = await client.get(url)  # без заголовка авторизации
    assert got.status_code == 200 and "Продукт" in got.content.decode("utf-8-sig")
    assert (await client.get("/api/downloads/unknown-token")).status_code == 404
